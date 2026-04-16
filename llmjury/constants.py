#
#    ADOBE CONFIDENTIAL
#    ___________________
#
#    Copyright 2025 Adobe
#    All Rights Reserved.
#
#    NOTICE:  All information contained herein is, and remains
#    the property of Adobe and its suppliers, if any. The intellectual
#    and technical concepts contained herein are proprietary to Adobe
#    and its suppliers and are protected by all applicable intellectual
#    property laws, including trade secret and copyright laws.
#    Dissemination of this information or reproduction of this material
#    is strictly forbidden unless prior written permission is obtained
#    from Adobe.
#
"""
Shared constants for LLM Jury evaluation.

This module contains configuration values and constants used across
the LLM Jury evaluation system.
"""

from enum import Enum

import openpyxl  # type: ignore[import-untyped]

# Evaluation Modes


class EvaluationMode(Enum):
    """Evaluation mode for LLM Jury."""

    METRICS = 'metrics'
    QUESTION_ANSWER = 'question_answer'
    COMPARISON = 'comparison'


# Prompt Directories
PROMPT_BASE_DIR = 'llmJURY'
CRITERIA_DEFINITION_BASE_DIR = 'criteriaDefinitionLLMJURY'

# Prompt Version Configuration Keys
PROMPT_CONFIG_BASE_VERSION_KEY = '_base_version_'
PROMPT_CONFIG_CUSTOMIZATION_KEYS_KEY = '_customization_keys_'

# Prompt Version Configuration File
PROMPT_VERSIONS_CONFIG_FILE = 'prompt_versions.json'


# Prompt IDs


class PromptFamily(Enum):
    """Base prompt family identifiers."""

    # Q&A prompts
    QNA_WITH_SOURCE = 'question_answer'
    QNA_WITHOUT_SOURCE = 'qna_without_source'

    # Comparison prompts
    COMPARISON_WITH_SOURCE = 'comparison_with_source'
    COMPARISON_WITHOUT_SOURCE = 'comparison_without_source'
    COMPARISON_WITH_GROUND_TRUTH = 'comparison_with_ground_truth'

    # Metrics prompts
    LLM_JURY = 'llm_jury'
    CONSOLIDATION = 'consolidation_prompt'

    # Custom
    CUSTOM = 'custom'


# Default values
DEFAULT_CONTENT_TYPE = 'summary'
DEFAULT_EVALUATION_MODE = 'metrics'

# Default prompt versions
DEFAULT_METRICS_VERSION = 'v2'
DEFAULT_COMPARISON_VERSION = 'v2'
DEFAULT_QUESTION_ANSWER_VERSION = 'v2'
DEFAULT_CONSOLIDATION_VERSION = 'v1'


# Sheet Configuration


class SheetName(Enum):
    """Available sheets (extend to add new sheets)."""

    MERGED = 'merged'
    REPORT = 'report'
    JURY_RESULTS = 'jury_results'
    INFO = 'info'
    EXECUTIVE_SUMMARY = 'executive_summary'


# Excel Formatting


class ExcelColor(Enum):
    """Excel fill colors for conditional formatting."""

    # Score/threshold colors
    RED = openpyxl.styles.PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    GREEN = openpyxl.styles.PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

    # Compression ratio tier colors
    MISTY_ROSE = openpyxl.styles.PatternFill(start_color='FFE4E1', end_color='FFE4E1', fill_type='solid')
    LAVENDER = openpyxl.styles.PatternFill(start_color='E6E6FA', end_color='E6E6FA', fill_type='solid')
    HONEYDEW = openpyxl.styles.PatternFill(start_color='F0FFF0', end_color='F0FFF0', fill_type='solid')


# Evaluation Thresholds
SCORE_THRESHOLD = 3.5
PERCENTAGE_THRESHOLD = 70
MAX_REASONABLE_SCORE = 100  # Max score cap to catch LLM errors (supports 1-5, 1-10, 1-100 scales)

# Column Naming Patterns (used to construct dynamic column names)
SCORE_SUFFIX = '_score'
REASONING_SUFFIX = '_reasoning'
PROMPT_VERSION_SUFFIX = '_prompt_version'
CALCULATED_SUFFIX = '_calculated_score'
STD_DEV_SUFFIX = '_std_deviation'
OVERALL_PREFIX = 'overall'

# Column name constants (derived from patterns)
OVERALL_SCORE = f'{OVERALL_PREFIX}{SCORE_SUFFIX}'  # 'overall_score'
OVERALL_REASONING = f'{OVERALL_PREFIX}{REASONING_SUFFIX}'  # 'overall_reasoning'

# Calculated columns
COMPRESSION_RATIO_COL = 'compression_ratio_percentage'

# LLM Jury Model
JURY_MODEL = 'jury'


# Evaluation Criteria


class EvaluationCriteria(Enum):
    """Available criteria for LLM evaluation."""

    ACCURACY = 'accuracy'
    BALANCE = 'balance'
    CLARITY = 'clarity'
    COHERENCE = 'coherence'
    COMPREHENSIVENESS = 'comprehensiveness'
    CONCISENESS = 'conciseness'
    NAVIGATION = 'navigation'
    TRUTHFULNESS = 'truthfulness'
    DUPLICACY = 'duplicacy'
    FLUENCY = 'fluency'
    METHOD_APPROPRIATENESS = 'method_appropriateness'
    DATA_SUFFICIENCY_RESOLUTION = 'data_sufficiency_resolution'
    QUERY_AMBIGUITY_RESOLUTION = 'query_ambiguity_resolution'
    ANALOGY = 'analogy'
    ENGAGEMENT = 'engagement'
    TONALITY = 'tonality'
    USEFULNESS = 'usefulness'
    HIERARCHICAL_STRUCTURE = 'hierarchical_structure'
    DEPTH_OF_REASONING = 'depth_of_reasoning'
    CUSTOM = 'custom'
    DIVERSITY = 'diversity'


# Metrics that should be evaluated without source text
METRICS_WITHOUT_SOURCE = [
    EvaluationCriteria.DUPLICACY.value,
    EvaluationCriteria.FLUENCY.value,
]
